"""
Gera as amostras de teste (tests/fixtures) recortando registros REAIS dos
arquivos oficiais. Nada é inventado: cabeçalhos e linhas são copiados.

Uso: python3 -I scripts/gerar-fixtures.py <pasta com DBFs EMOP> <SINAPI_Referência_AAAA_MM.xlsx> <saída>
"""
import os, re, struct, sys, zipfile
import openpyxl

pasta_emop, xlsx_sinapi, saida = sys.argv[1:4]
os.makedirs(saida, exist_ok=True)

# ----------------------------------------------------------------- EMOP (DBF)
def ler_dbf(caminho):
    b = open(caminho, "rb").read()
    n, hl, rl = struct.unpack("<IHH", b[4:12])
    campos, o = [], 32
    while b[o] != 0x0D:
        nome = b[o:o + 11].split(b"\0")[0].decode("latin1"); tam = b[o + 16]
        campos.append((nome, tam)); o += 32
    regs = []
    for i in range(n):
        r = b[hl + i * rl: hl + (i + 1) * rl]
        d, p = {}, 1
        for nome, tam in campos:
            d[nome] = r[p:p + tam].decode("cp850").strip(); p += tam
        regs.append((r, d))
    return b[:hl], regs

def gravar_dbf(caminho, cab, regs):
    cab = bytearray(cab); cab[4:8] = struct.pack("<I", len(regs))
    with open(caminho, "wb") as f:
        f.write(cab); [f.write(r) for r in regs]; f.write(b"\x1a")

mmaa = re.search(r"(\d{4})\.dbf$", [f for f in os.listdir(pasta_emop) if re.match(r"DIPR\d{4}\.dbf$", f, re.I)][0], re.I).group(1)
t = {n: ler_dbf(os.path.join(pasta_emop, f"{n}{mmaa}.dbf")) for n in ["EMOP", "DIPR", "COMP", "MAT", "ELEM", "REUT"]}

# serviços escolhidos: normais (0/A), com PERCENTUAL, um que usa elementar reutilizado, e um índice setorial
reut_elem = {d["ELEMENTAR"] for _, d in t["REUT"][1]}
usa_reut = next(d["CODIGO"] for _, d in t["COMP"][1] if d["ELEMENTAR"] in reut_elem and d["CODIGO"].endswith("-0"))
com_pct = next(d["CODIGO"] for _, d in t["COMP"][1] if float(d["PERCENTUAL"] or 0) > 0 and d["CODIGO"].endswith("-0"))
base = {"01.001.0001", "01.001.0002", com_pct[:-2], usa_reut[:-2], "05.100.9999"}
servicos = {d["CODIGO"] for _, d in t["DIPR"][1] if d["CODIGO"][:-2] in base}
comp = [(r, d) for r, d in t["COMP"][1] if d["CODIGO"] in servicos]
elems = {d["ELEMENTAR"] for _, d in comp} | {"01294"}  # 01294: publicado sem unidade no REUT
gravar_dbf(f"{saida}/DIPR{mmaa}.dbf", t["DIPR"][0], [r for r, d in t["DIPR"][1] if d["CODIGO"] in servicos])
gravar_dbf(f"{saida}/EMOP{mmaa}.dbf", t["EMOP"][0], [r for r, d in t["EMOP"][1] if d["CODIGO"] in servicos])
gravar_dbf(f"{saida}/COMP{mmaa}.dbf", t["COMP"][0], [r for r, _ in comp])
gravar_dbf(f"{saida}/MAT{mmaa}.dbf", t["MAT"][0], [r for r, d in t["MAT"][1] if d["ELEMENTAR"] in elems])
gravar_dbf(f"{saida}/REUT{mmaa}.dbf", t["REUT"][0], [r for r, d in t["REUT"][1] if d["ELEMENTAR"] in elems])
gravar_dbf(f"{saida}/ELEM{mmaa}.dbf", t["ELEM"][0], [r for r, d in t["ELEM"][1] if d["ELEMENTAR"] in elems])
print("EMOP: serviços", sorted(servicos), "elementares", len(elems))

# ------------------------------------------------------------- SINAPI (XLSX)
wbv = openpyxl.load_workbook(xlsx_sinapi, read_only=True, data_only=True)
wbf = openpyxl.load_workbook(xlsx_sinapi, read_only=True, data_only=False)

def linhas(aba, formulas=False):
    return list((wbf if formulas else wbv)[aba].iter_rows(values_only=True))

ana = linhas("Analítico")
# composições: a 1ª do arquivo (tem subcomposição), sua subcomposição, e uma com custo zerado no RJ
csd_v, csd_f = linhas("CSD"), linhas("CSD", True)
cod = lambda f: re.search(r"MATCH\((\d+),", f).group(1) if isinstance(f, str) else str(f)
zerada = next(cod(csd_f[i][1]) for i in range(10, len(csd_v)) if csd_v[i][40] in (0, None))
primeira = str(ana[10][1])
sub = next(str(r[3]) for r in ana[10:] if str(r[1]) == primeira and r[2] == "COMPOSICAO")
comps = {primeira, sub, zerada}
ana_sel = [r for r in ana[:10]] + [r for r in ana[10:] if str(r[1]) in comps]
insumos = {str(r[3]) for r in ana_sel[10:] if r[2] == "INSUMO"}
isd = linhas("ISD")
sem_preco_rj = next(str(r[1]) for r in isd[10:] if r[23] is None)
insumos.add(sem_preco_rj)

wb = openpyxl.Workbook(); wb.remove(wb.active)
for aba in ["ISD", "ICD", "ISE"]:
    ws = wb.create_sheet(aba); src = linhas(aba)
    for r in src[:10] + [r for r in src[10:] if str(r[1]) in insumos]: ws.append(list(r))
for aba in ["CSD", "CCD", "CSE"]:
    ws = wb.create_sheet(aba); v, f = linhas(aba), linhas(aba, True)
    for i, r in enumerate(v):
        if i < 10 or cod(f[i][1]) in comps:
            row = list(r)
            if i >= 10: row[1] = f[i][1]  # mantém a fórmula original (valor gravado no oficial é 0)
            ws.append(row)
ws = wb.create_sheet("Analítico")
for r in ana_sel: ws.append(list(r))
nome = os.path.basename(xlsx_sinapi)
tmp = os.path.join(saida, nome); wb.save(tmp)
ano_mes = re.search(r"(\d{4})_(\d{2})", nome).groups()
with zipfile.ZipFile(os.path.join(saida, f"SINAPI-{ano_mes[0]}-{ano_mes[1]}-amostra.zip"), "w", zipfile.ZIP_DEFLATED) as z:
    z.write(tmp, nome)
os.remove(tmp)
print("SINAPI: composições", sorted(comps), "insumos", len(insumos), "sem preço RJ", sem_preco_rj)
